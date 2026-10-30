"""Convert analysed pages (typed blocks) into editable formats, keeping structure:
headings, paragraphs (alignment, size, bold), tables, images, captions.
"""
from __future__ import annotations

import base64
import html
import io
import re
import time
import uuid
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

TAMIL_FONT = "Nirmala UI"   # ships with Windows 8+; Word falls back gracefully elsewhere
LATIN_FONT = "Calibri"
HTML_LANG = {"ta": "ta", "en": "en", "hi": "hi", "ml": "ml", "te": "te", "kn": "kn", "bn": "bn",
             "gu": "gu", "pa": "pa", "or": "or", "si": "si", "ar": "ar"}

DEFAULTS = {"images": True, "headers": False, "line_breaks": False, "page_breaks": True}


# ---------------------------------------------------------------------------
# shared helpers
# ---------------------------------------------------------------------------
def opts(**kw) -> dict:
    o = dict(DEFAULTS)
    o.update({k: v for k, v in kw.items() if v is not None})
    return o


def blocks_of(page: dict, o: dict) -> list[dict]:
    out = []
    for b in page.get("blocks", []):
        t = b["type"]
        if t == "ignore" or (t in ("header", "footer") and not o["headers"]):
            continue
        if t == "image" and not o["images"]:
            continue
        out.append(b)
    return out


_HYPHEN = re.compile(r"(\w)-\n(\w)")


def block_text(b: dict, o: dict) -> str:
    t = b.get("text", "")
    if o["line_breaks"] or b["type"].startswith("heading"):
        return t
    t = _HYPHEN.sub(r"\1\2", t)          # re-join words hyphenated at line end
    return re.sub(r"\s*\n\s*", " ", t)


def main_lang(b) -> str:
    return HTML_LANG.get((b.get("lang") or "").split("+")[0], "")


def image_width_in(page: dict, b: dict) -> float:
    dpi = page.get("dpi", 300) * page.get("scale", 1.0)
    return max(0.8, min(6.3, b["bbox"][2] / dpi))


# ---------------------------------------------------------------------------
# Plain text / Markdown
# ---------------------------------------------------------------------------
def to_txt(title, pages, fig, o) -> bytes:
    out = []
    for p in pages:
        if o["page_breaks"]:
            out.append(f"===== Page {p['page']} =====")
        for b in blocks_of(p, o):
            if b["type"] == "table":
                out.append("\n".join("\t".join(c.replace("\n", " ") for c in r) for r in b["rows"]))
            elif b["type"] == "image":
                out.append("[Image]")
            else:
                out.append(block_text(b, o))
            out.append("")
    return "\n".join(out).encode("utf-8")


def to_md(title, pages, fig, o) -> bytes:
    """Markdown + images folder, zipped."""
    buf = io.BytesIO()
    lines = [f"# {title}", ""]
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in pages:
            if o["page_breaks"]:
                lines += [f"<!-- page {p['page']} -->", ""]
            for b in blocks_of(p, o):
                t = b["type"]
                if t.startswith("heading"):
                    lines.append("#" * (int(t[-1]) + 1) + " " + b["text"].replace("\n", " "))
                elif t == "table":
                    rows = [[c.replace("|", "\\|").replace("\n", "<br>") for c in r] for r in b["rows"]]
                    lines.append("| " + " | ".join(rows[0]) + " |")
                    lines.append("|" + "---|" * len(rows[0]))
                    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                elif t == "image":
                    name = f"images/{b['image']}"
                    z.writestr(name, fig(p, b))
                    lines.append(f"![]({name})")
                elif t == "caption":
                    lines.append(f"*{block_text(b, o)}*")
                else:
                    txt = block_text(b, o)
                    lines.append(txt.replace("\n", "  \n") if o["line_breaks"] else txt)
                lines.append("")
        z.writestr("book.md", "\n".join(lines))
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def _fonts(run, size=None, bold=None, italic=None, color=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi"):
        rf.set(qn(a), LATIN_FONT)
    rf.set(qn("w:cs"), TAMIL_FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:bidi"), "ta-IN")
    if size:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rpr.append(szcs)
    if bold is not None:
        run.font.bold = bold
        if bold:
            rpr.append(OxmlElement("w:bCs"))
    if italic:
        run.font.italic = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _runs(par, text, **style):
    parts = text.split("\n")
    for i, line in enumerate(parts):
        r = par.add_run(line)
        _fonts(r, **style)
        if i < len(parts) - 1:
            r.add_break()


ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def to_docx(title, pages, fig, o) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        stl = doc.styles[name]
        rpr = stl.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        rf.set(qn("w:cs"), TAMIL_FONT)
    doc.core_properties.title = title

    body_pts = [b["style"]["size"] for p in pages for b in p.get("blocks", [])
                if b["type"] == "paragraph" and b.get("style")]
    body_pt = sorted(body_pts)[len(body_pts) // 2] if body_pts else 11
    k = 11 / body_pt if body_pt else 1  # normalise sizes so body text is 11pt in Word

    first = True
    for p in pages:
        if not first and o["page_breaks"]:
            doc.add_page_break()
        first = False
        for b in blocks_of(p, o):
            t, stl = b["type"], b.get("style") or {}
            size = round(stl.get("size", body_pt) * k * 2) / 2 if stl.get("size") else None
            if t.startswith("heading"):
                h = doc.add_heading(level=int(t[-1]))
                _runs(h, b["text"], size=size)
                h.alignment = ALIGN.get(stl.get("align"), None)
            elif t == "table":
                rows = b["rows"]
                ncol = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=ncol)
                tbl.style = "Table Grid"
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, r in enumerate(rows):
                    for j in range(ncol):
                        cell = tbl.cell(i, j)
                        cell.paragraphs[0].text = ""
                        _runs(cell.paragraphs[0], r[j] if j < len(r) else "", size=10, bold=(i == 0))
                doc.add_paragraph()
            elif t == "image":
                doc.add_picture(io.BytesIO(fig(p, b)), width=Inches(image_width_in(p, b)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                par = doc.add_paragraph()
                italic = t == "caption"
                color = "777777" if t in ("header", "footer") else None
                _runs(par, block_text(b, o), size=min(size or 11, 11) if t != "paragraph" else size,
                      bold=stl.get("bold") or None, italic=italic, color=color)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if t == "caption" else ALIGN.get(stl.get("align"))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def to_xlsx(title, pages, fig, o) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Content"
    hdr_font, body_font = Font(name=TAMIL_FONT, bold=True), Font(name=TAMIL_FONT)
    fill = PatternFill("solid", fgColor="F3E7DD")
    ws.append(["Page", "Type", "Language", "Text"])
    for c in ws[1]:
        c.font, c.fill = hdr_font, fill
    n_tab = 0
    for p in pages:
        for b in blocks_of(p, o):
            if b["type"] == "image":
                continue
            if b["type"] == "table":
                n_tab += 1
                name = f"P{p['page']} Table {n_tab}"[:31]
                ws.append([p["page"], "table", b.get("lang", ""), f"→ sheet '{name}'"])
                ts = wb.create_sheet(name)
                for i, r in enumerate(b["rows"]):
                    ts.append([_num(c) for c in r])
                    for c in ts[i + 1]:
                        c.font = hdr_font if i == 0 else body_font
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                        if i == 0:
                            c.fill = fill
                for j in range(1, max(len(r) for r in b["rows"]) + 1):
                    width = max(len(str(r[j - 1])) if j - 1 < len(r) else 0 for r in b["rows"])
                    ts.column_dimensions[get_column_letter(j)].width = min(60, max(10, width * 1.3))
                ts.freeze_panes = "A2"
            else:
                ws.append([p["page"], b["type"], b.get("lang", ""), block_text(b, o)])
                ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(ws.max_row, 4).font = body_font
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 7, 11
    ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 10, 100
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_NUM = re.compile(r"^-?[\d,]+(\.\d+)?$")


def _num(s: str):
    """Keep numbers as numbers in Excel (1,234.5 -> 1234.5)."""
    t = s.strip()
    if _NUM.match(t) and any(ch.isdigit() for ch in t):
        try:
            v = float(t.replace(",", ""))
            return int(v) if v.is_integer() and "." not in t else v
        except ValueError:
            pass
    return s


# ---------------------------------------------------------------------------
# HTML / EPUB
# ---------------------------------------------------------------------------
CSS = """
body{font-family:"Noto Serif Tamil","Noto Sans Tamil Supplement",Georgia,serif;max-width:46em;margin:2em auto;padding:0 1em;line-height:1.75;color:#1f1b16}
h1,h2,h3{font-family:"Noto Sans Tamil","Noto Sans Tamil Supplement",sans-serif;line-height:1.35}
table{border-collapse:collapse;margin:1em auto;font-size:.95em}
td,th{border:1px solid #bbb;padding:.35em .6em;vertical-align:top;text-align:left}
th{background:#f3ece2}
figure{margin:1.2em 0;text-align:center} img{max-width:100%;height:auto}
.caption{text-align:center;font-style:italic;font-size:.9em;color:#555}
.hf{color:#888;font-size:.85em}
.page{border-top:1px dashed #ccc;margin:2.5em 0 1em;font:12px sans-serif;color:#999}
"""


def _html_blocks(p, o, img_src) -> str:
    out = []
    for b in blocks_of(p, o):
        t, stl = b["type"], b.get("style") or {}
        lang = main_lang(b)
        la = f' lang="{lang}"' if lang else ""
        style = []
        if stl.get("align") and stl["align"] != "left" and t != "caption":
            style.append(f"text-align:{stl['align']}")
        if stl.get("bold") and t == "paragraph":
            style.append("font-weight:600")
        st_attr = f' style="{";".join(style)}"' if style else ""
        txt = html.escape(block_text(b, o)).replace("\n", "<br>\n")
        if t.startswith("heading"):
            n = int(t[-1])
            out.append(f"<h{n}{la}{st_attr}>{txt}</h{n}>")
        elif t == "table":
            rows = b["rows"]
            head = "".join(f"<th>{html.escape(c).replace(chr(10), '<br>')}</th>" for c in rows[0])
            body = "".join("<tr>" + "".join(f"<td>{html.escape(c).replace(chr(10), '<br>')}</td>" for c in r) + "</tr>"
                           for r in rows[1:])
            out.append(f"<table{la}><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        elif t == "image":
            out.append(f'<figure><img src="{img_src(p, b)}" alt=""/></figure>')
        elif t == "caption":
            out.append(f'<p class="caption"{la}>{txt}</p>')
        elif t in ("header", "footer"):
            out.append(f'<p class="hf"{la}>{txt}</p>')
        else:
            out.append(f"<p{la}{st_attr}>{txt}</p>")
    return "\n".join(out)


def to_html(title, pages, fig, o) -> bytes:
    src = lambda p, b: "data:image/jpeg;base64," + base64.b64encode(fig(p, b)).decode()
    parts = []
    for p in pages:
        if o["page_breaks"]:
            parts.append(f'<div class="page">Page {p["page"]}</div>')
        parts.append(_html_blocks(p, o, src))
    doc = f"""<!doctype html>
<html lang="ta"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+Tamil:wght@400;600&family=Noto+Serif+Tamil&family=Noto+Sans+Tamil+Supplement&display=swap" rel="stylesheet">
<style>{CSS}</style></head><body>
<h1>{html.escape(title)}</h1>
{chr(10).join(parts)}
</body></html>"""
    return doc.encode("utf-8")


def to_epub(title, pages, fig, o) -> bytes:
    """EPUB 3 e-book: one chapter per page, images included."""
    buf = io.BytesIO()
    uid = f"urn:uuid:{uuid.uuid4()}"
    images, chapters = [], []
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>""")

        def src(p, b):
            name = f"images/{b['image']}"
            if name not in images:
                z.writestr("OEBPS/" + name, fig(p, b), compress_type=zipfile.ZIP_DEFLATED)
                images.append(name)
            return name

        for p in pages:
            body = _html_blocks(p, dict(o, page_breaks=False), src)
            name = f"page{p['page']:04d}.xhtml"
            chapters.append((name, p["page"]))
            z.writestr("OEBPS/" + name, f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" xml:lang="ta" lang="ta"><head><meta charset="utf-8"/>
<title>{html.escape(title)} – {p['page']}</title><link rel="stylesheet" href="style.css"/></head>
<body>{_xhtml(body)}</body></html>""", compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/style.css", CSS)
        nav = "".join(f'<li><a href="{n}">Page {pg}</a></li>' for n, pg in chapters)
        z.writestr("OEBPS/nav.xhtml", f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head>
<body><nav epub:type="toc"><h1>Contents</h1><ol>{nav}</ol></nav></body></html>""")
        manifest = "".join(f'<item id="c{i}" href="{n}" media-type="application/xhtml+xml"/>' for i, (n, _) in enumerate(chapters))
        manifest += "".join(f'<item id="i{i}" href="{n}" media-type="image/jpeg"/>' for i, n in enumerate(images))
        spine = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
        z.writestr("OEBPS/content.opf", f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="bookid">{uid}</dc:identifier>
<dc:title>{html.escape(title)}</dc:title><dc:language>ta</dc:language>
<meta property="dcterms:modified">{time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())}</meta></metadata>
<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
<item id="css" href="style.css" media-type="text/css"/>{manifest}</manifest>
<spine>{spine}</spine></package>""")
    return buf.getvalue()


def _xhtml(s: str) -> str:
    return s.replace("<br>", "<br/>")


# ---------------------------------------------------------------------------
FORMATS = {
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "html": (to_html, "text/html; charset=utf-8", "html"),
    "epub": (to_epub, "application/epub+zip", "epub"),
    "md": (to_md, "application/zip", "md.zip"),
    "txt": (to_txt, "text/plain; charset=utf-8", "txt"),
}
